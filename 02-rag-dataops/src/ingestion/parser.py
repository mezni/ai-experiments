from __future__ import annotations

import html as html_module
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pydantic import BaseModel, Field
from pypdf import PdfReader


class ParseError(Exception):
    """Raised when a single document cannot be parsed."""


class ParsedDocument(BaseModel):
    """Normalized document produced by a parser (PROJECT.md §5)."""

    document_id: str
    source: str
    format: str
    text: str = Field(min_length=0)
    metadata: dict[str, Any] = Field(default_factory=dict)
    collected_at: str = ""


def _now() -> str:
    return datetime.now(UTC).isoformat()


def parse_pdf(path: Path, document_id: str, source: str = "filesystem") -> ParsedDocument:
    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise ParseError(f"PDF parsing failed for {path}: {exc}") from exc
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")  # per-page best effort
    text = "\n\n".join(page.strip() for page in pages if page.strip())
    return ParsedDocument(
        document_id=document_id,
        source=source,
        format="pdf",
        text=text,
        metadata={"pages": len(reader.pages), "path": str(path)},
        collected_at=_now(),
    )


def parse_docx(path: Path, document_id: str, source: str = "filesystem") -> ParsedDocument:
    try:
        doc = DocxDocument(str(path))
    except Exception as exc:
        raise ParseError(f"DOCX parsing failed for {path}: {exc}") from exc
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    text = "\n\n".join(paragraphs)
    return ParsedDocument(
        document_id=document_id,
        source=source,
        format="docx",
        text=text,
        metadata={"paragraphs": len(paragraphs), "path": str(path)},
        collected_at=_now(),
    )


def _extract_headings(content: str) -> list[str]:
    soup = BeautifulSoup(content, "html.parser")
    return [tag.get_text(" ", strip=True) for tag in soup.find_all(["h1", "h2", "h3"])]


def parse_html(
    content: str, document_id: str, source: str = "html", url: str | None = None
) -> ParsedDocument:
    soup = BeautifulSoup(content, "html.parser")
    for tag in ["script", "style", "nav", "footer"]:
        for element in soup.find_all(tag):
            element.decompose()
    text = soup.get_text("\n", strip=True)
    metadata: dict[str, Any] = {"headings": _extract_headings(content)}
    if url:
        metadata["url"] = url
    return ParsedDocument(
        document_id=document_id,
        source=source,
        format="html",
        text=text,
        metadata=metadata,
        collected_at=_now(),
    )


def parse_text_file(path: Path, document_id: str, source: str = "filesystem") -> ParsedDocument:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        raise ParseError(f"TXT parsing failed for {path}: {exc}") from exc
    return ParsedDocument(
        document_id=document_id,
        source=source,
        format="txt",
        text=text,
        metadata={"path": str(path)},
        collected_at=_now(),
    )


def parse_database_row(
    *,
    table: str,
    primary_key: str,
    row: dict[str, Any],
    projection_columns: list[str] | None = None,
    source: str = "database",
    document_id: str | None = None,
) -> ParsedDocument:
    """Project one database row into a text document (PROJECT.md §5)."""
    columns = projection_columns or sorted(row)
    header = f"Table {table}  Row {primary_key} = {row[primary_key]}"
    lines = [f"{column}: {row.get(column, '')}" for column in columns if column in row]
    text = header + "\n" + "\n".join(lines)
    doc_id = document_id or f"{table}/{row[primary_key]}"
    return ParsedDocument(
        document_id=doc_id,
        source=source,
        format="database_row",
        text=text,
        metadata={
            "table": table,
            "primary_key": primary_key,
            "row": {col: row[col] for col in columns if col in row},
        },
        collected_at=_now(),
    )


def escape_html(text: str) -> str:
    return html_module.escape(text)


_EXTENSION_ROUTING = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".html": lambda path, document_id, source="filesystem": parse_html(
        path.read_text(encoding="utf-8", errors="replace"),
        document_id,
        source=source,
        url=path.as_uri(),
    ),
    ".htm": lambda path, document_id, source="filesystem": parse_html(
        path.read_text(encoding="utf-8", errors="replace"),
        document_id,
        source=source,
        url=path.as_uri(),
    ),
    ".txt": parse_text_file,
}


def parse_file(path: Path, document_id: str, source: str = "filesystem") -> ParsedDocument:
    """Route a discovered file to its parser by extension."""
    parser = _EXTENSION_ROUTING.get(path.suffix.lower())
    if parser is None:
        raise ParseError(f"No parser registered for extension {path.suffix!r} ({path})")
    return parser(path, document_id, source)